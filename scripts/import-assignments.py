#!/usr/bin/env python3
"""Import Bealls SOC station assignments from Excel into src/app/data/soc-data.json."""
from __future__ import annotations

from openpyxl import load_workbook
from pathlib import Path
import json
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path.home() / "Downloads" / "Store Assignments Dataminr 3-22-26 (Ty Copy).xlsx"
OUT = ROOT / "src" / "app" / "data" / "soc-data.json"
BEALLS_PATH = ROOT / "src" / "app" / "data" / "bealls-locations.json"


def load_open_store_ids() -> set[str]:
    if not BEALLS_PATH.exists():
        raise SystemExit(f"Run npm run fetch-stores first — missing {BEALLS_PATH}")
    data = json.loads(BEALLS_PATH.read_text())
    return {
        pad_id(sid)
        for sid, store in data.get("stores", {}).items()
        if not store.get("closed")
    }


def load_bealls_labels() -> dict[str, str]:
    data = json.loads(BEALLS_PATH.read_text())
    return data.get("storeLabels", {})


def load_bealls_coords() -> dict[str, list[float]]:
    data = json.loads(BEALLS_PATH.read_text())
    return data.get("storeCoords", {})


def keep_store(sid: str | None, open_ids: set[str]) -> bool:
    return bool(sid and sid in open_ids)

CITY_COORDS = {
    ("Englewood", "FL"): [26.962, -82.352], ("Largo", "FL"): [27.909, -82.787],
    ("Sarasota", "FL"): [27.336, -82.531], ("Pt St Lucie", "FL"): [27.273, -80.358],
    ("Bradenton", "FL"): [27.499, -82.574], ("Vero Beach", "FL"): [27.638, -80.397],
    ("N Ft Myers", "FL"): [26.677, -81.898], ("Ormond Beach", "FL"): [29.286, -81.056],
    ("St. Augustine", "FL"): [29.901, -81.312], ("Boynton Beach", "FL"): [26.525, -80.066],
    ("Riverview", "FL"): [27.866, -82.326], ("Sanford", "FL"): [28.803, -81.269],
    ("Cortez", "FL"): [27.469, -82.686], ("Auburndale", "FL"): [28.065, -81.789],
    ("Daytona", "FL"): [29.211, -81.022], ("Gun Barrel City", "TX"): [32.334, -96.151],
    ("Orangeburg", "SC"): [33.491, -80.856], ("Jacksonville", "IL"): [39.734, -90.229],
    ("Lebanon", "TN"): [36.208, -86.291], ("Ocala", "FL"): [29.187, -82.140],
    ("Franklin", "NC"): [35.182, -83.382], ("Bartow", "FL"): [27.896, -81.843],
    ("Key Largo", "FL"): [25.087, -80.447], ("Pecos", "TX"): [31.423, -103.493],
    ("Lakeland", "FL"): [28.040, -81.949], ("Mt Airy", "NC"): [36.499, -80.607],
    ("Monahans", "TX"): [31.594, -102.893], ("Shelby", "NC"): [35.292, -81.535],
    ("Waynesville", "NC"): [35.488, -82.987], ("Granite Falls", "NC"): [35.796, -81.431],
    ("Hillsboro", "FL"): [29.387, -82.137], ("Ft. Pierce", "FL"): [27.447, -80.326],
    ("Tullahoma", "TN"): [35.362, -86.209], ("High Point", "NC"): [35.956, -80.005],
    ("Washington", "NC"): [35.547, -77.052], ("Sulphur Springs", "TX"): [33.138, -95.601],
}


def pad_id(n):
    return str(int(n)).zfill(3)


def parse_label(label):
    if not label:
        return {"city": "Unknown", "state": ""}
    m = re.search(r"-\s*(.+)$", str(label))
    text = m.group(1).strip() if m else str(label)
    text = text.replace("beall's", "").replace("bealls", "").replace("Bealls", "").strip()
    parts = text.rsplit(" ", 1)
    if len(parts) == 2 and len(parts[1]) == 2:
        return {"city": parts[0].strip(), "state": parts[1]}
    return {"city": text, "state": ""}


def coord_for_label(label):
    loc = parse_label(label)
    key = (loc["city"], loc["state"])
    if key in CITY_COORDS:
        return CITY_COORDS[key]
    for (city, st), coord in CITY_COORDS.items():
        if city.lower() in loc["city"].lower() or loc["city"].lower() in city.lower():
            return coord
    return None


def extract_exts(cell) -> list[str]:
    """Pull extension numbers from cells like 'ext. 2253', 'ext 3592', 'ext. 1350 or 3501'."""
    if cell is None:
        return []
    text = str(cell)
    return [m.group(1) for m in re.finditer(r"(?<!\d)(\d{3,4})(?!\d)", text)]


def soc_site_from_ext(ext: str) -> str | None:
    """Fallback FL/TX tagging by known SOC extension ranges."""
    try:
        n = int(str(ext).strip())
    except (TypeError, ValueError):
        return None
    if n in (2253, 2263) or 3580 <= n <= 3593:
        return "TX"
    if n == 1350 or 3501 <= n <= 3525:
        return "FL"
    return None


def parse_soc_contacts(wb) -> dict[str, dict[str, str]]:
    """
    Map extension → 'FL' | 'TX' from the SOC Contacts sheet.
    TX names/exts are in cols B/D; FL names/exts in cols F/G.
    """
    if "SOC Contacts" not in wb.sheetnames:
        return {"by_ext": {}, "by_name": {}}
    ws = wb["SOC Contacts"]
    by_ext: dict[str, str] = {}
    by_name: dict[str, str] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 4:
            continue
        # TX block: name col 1, ext col 3
        tx_name = row[1] if len(row) > 1 else None
        tx_exts = extract_exts(row[3] if len(row) > 3 else None)
        if isinstance(tx_name, str) and tx_name.strip() and tx_name.strip().upper() not in ("OPEN", "N/A"):
            clean = tx_name.strip()
            for ext in tx_exts:
                by_ext[ext] = "TX"
                by_name[clean.lower()] = "TX"
        else:
            for ext in tx_exts:
                by_ext[ext] = "TX"

        # FL block: name col 5, ext col 6
        fl_name = row[5] if len(row) > 5 else None
        fl_exts = extract_exts(row[6] if len(row) > 6 else None)
        if isinstance(fl_name, str) and fl_name.strip() and fl_name.strip().upper() not in ("OPEN", "N/A"):
            clean = fl_name.strip()
            for ext in fl_exts:
                by_ext[ext] = "FL"
                by_name[clean.lower()] = "FL"
        else:
            for ext in fl_exts:
                by_ext[ext] = "FL"

    return {"by_ext": by_ext, "by_name": by_name}


def resolve_soc_site(ext: str, name: str, contacts: dict[str, dict[str, str]]) -> str | None:
    by_ext = contacts.get("by_ext", {})
    by_name = contacts.get("by_name", {})
    if ext in by_ext:
        return by_ext[ext]
    # Name match (contacts often include role prefixes)
    name_l = name.lower().strip()
    if name_l in by_name:
        return by_name[name_l]
    for contact_name, site in by_name.items():
        if name_l in contact_name or contact_name in name_l:
            return site
    return soc_site_from_ext(ext)


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        raise SystemExit(f"Excel file not found: {xlsx}")

    open_ids = load_open_store_ids()
    bealls_labels = load_bealls_labels()
    bealls_coords = load_bealls_coords()

    wb = load_workbook(xlsx, data_only=True, read_only=True)
    soc_contacts = parse_soc_contacts(wb)

    ws = wb["Stores"]
    rows = list(ws.iter_rows(values_only=True))
    store_slots = {}
    for i, row in enumerate(rows):
        if i < 2 or not isinstance(row[0], (int, float)):
            continue
        sid = pad_id(int(row[0]))
        if sid not in open_ids:
            continue
        slots = {
            "before9": pad_id(row[1]) if isinstance(row[1], (int, float)) and row[1] > 0 else None,
            "after9": pad_id(row[2]) if isinstance(row[2], (int, float)) and row[2] > 0 else None,
            "after10": pad_id(row[4]) if isinstance(row[4], (int, float)) and row[4] > 0 else None,
            "after11": pad_id(row[6]) if isinstance(row[6], (int, float)) and row[6] > 0 else None,
        }
        store_slots[sid] = {
            k: v if keep_store(v, open_ids) else None for k, v in slots.items()
        }

    ws = wb["Assignments"]
    focus_blocks = []
    all_hot = []
    top_to_focus = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 3:
            continue
        focus = [pad_id(int(c)) for c in row[4:17] if isinstance(c, (int, float)) and c > 0 and pad_id(int(c)) in open_ids]
        if not focus:
            break
        hot_raw = pad_id(int(row[17])) if len(row) > 17 and isinstance(row[17], (int, float)) and row[17] > 0 else None
        hot = hot_raw if keep_store(hot_raw, open_ids) else None
        block = {"focusStores": focus, "hotStore": hot, "dataminrOp": row[22] if len(row) > 22 else None}
        focus_blocks.append(block)
        for s in focus:
            top_to_focus.setdefault(s, focus)
        if hot:
            all_hot.append(hot)

    ws = wb["Top_Stores"]
    operators = []
    labels = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        ext = row[1]
        name = row[0] if isinstance(row[0], str) and row[0].strip() else None
        # Some rows put the name in col 3 when Operator (col 0) is blank
        if not name and len(row) > 3 and isinstance(row[3], str) and row[3].strip():
            name = row[3].strip()
        if ext is None or not name:
            continue
        if not isinstance(ext, (int, float, str)):
            continue
        ext = str(int(ext)) if isinstance(ext, float) else str(ext).strip()
        chain = [pad_id(c) for c in (row[12:26] if len(row) > 12 else []) if isinstance(c, (int, float)) and 0 < c < 9000 and pad_id(int(c)) in open_ids]
        top = int(row[9]) if len(row) > 9 and isinstance(row[9], (int, float)) else None
        top_label = row[10] if len(row) > 10 else None
        top_sid = pad_id(top) if top and keep_store(pad_id(top), open_ids) else None
        if top_sid:
            labels[top_sid] = bealls_labels.get(top_sid) or (str(top_label).strip() if top_label else f"Store {top_sid}")
        extra = [pad_id(int(c)) for c in row[3:5] if isinstance(c, (int, float)) and c > 0 and pad_id(int(c)) in open_ids]
        hot_raw = pad_id(int(row[26])) if len(row) > 26 and isinstance(row[26], (int, float)) and row[26] > 0 else None
        hot = hot_raw if keep_store(hot_raw, open_ids) else None
        store_count = int(row[8]) if len(row) > 8 and isinstance(row[8], (int, float)) else len(chain)
        # Top_Stores chain (cols 12–25). If Excel formulas weren't cached, fall back to
        # the Assignments "Focus Stores" block that contains this operator's top store.
        assigned = list(chain)
        if not assigned and top_sid and top_sid in top_to_focus:
            assigned = list(top_to_focus[top_sid])
        if not assigned and top_sid:
            assigned = [top_sid]
        # Cols 3–4 often hold extra/hot-adjacent stores — merge without dupes
        for e in extra:
            if e not in assigned:
                assigned.append(e)
        if hot and hot not in assigned:
            # keep hot separate; don't force into assigned unless it was in chain/focus
            pass
        site = resolve_soc_site(str(ext), name.strip(), soc_contacts)
        operators.append({
            "id": f"op-{ext}",
            "name": name.strip(),
            "extension": str(ext),
            "socSite": site,
            "storeCount": store_count,
            "topStore": top_sid,
            "topLabel": labels.get(top_sid) if top_sid else None,
            "assignedStores": assigned,
            "hotStore": hot,
            "extraStores": extra,
        })
   
    for op in operators:
        if op.get("topStore") and op.get("topLabel"):
            labels[op["topStore"]] = op["topLabel"]

    labels = {sid: bealls_labels[sid] for sid in labels if sid in open_ids and sid in bealls_labels}

    ws = wb["Dataminr Alerts"]
    dataminr = []
    watch_from_dm = set()
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < 4:
            continue
        store = row[1]
        if not isinstance(store, (int, float)):
            continue
        sid = pad_id(int(store))
        if sid not in open_ids:
            continue
        watch_from_dm.add(sid)
        dataminr.append({
            "store": sid,
            "activity": row[2],
            "start": row[3].isoformat() if hasattr(row[3], "isoformat") else str(row[3]),
            "end": row[4].isoformat() if hasattr(row[4], "isoformat") else str(row[4]),
            "comments": row[5],
            "operator": row[6],
        })

    priority_ids = [s for s in (focus_blocks[0]["focusStores"] if focus_blocks else []) if s in open_ids]
    opportunity_ids = [s for s in dict.fromkeys(all_hot) if s in open_ids]
    watch_ids = [s for s in dict.fromkeys(all_hot + list(watch_from_dm)) if s in open_ids]
    watch_list = [{
        "id": sid,
        "city": parse_label(bealls_labels.get(sid, ""))["city"] if bealls_labels.get(sid) else f"Store {sid}",
        "region": parse_label(bealls_labels.get(sid, ""))["state"] or "—",
        "reason": "Hot store or active Dataminr alert",
        "tier": "—",
        "socMonitored": True,
        "lastReview": "—",
        "flag": "Hot Store" if sid in all_hot else "Dataminr",
    } for sid in watch_ids]

    store_coords = {sid: bealls_coords[sid] for sid in open_ids if sid in bealls_coords}

    out = {
        "sourceFile": xlsx.name,
        "importedAt": "2026-03-22",
        "operators": operators,
        "focusBlocks": focus_blocks,
        "priorityStoreIds": priority_ids,
        "opportunityStoreIds": opportunity_ids,
        "watchList": watch_list,
        "storeLabels": labels,
        "storeSlots": store_slots,
        "storeCoords": store_coords,
        "dataminrAlerts": dataminr,
        "defaultOperatorExt": "3516",
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, indent=2))
    fl_n = sum(1 for o in operators if o.get("socSite") == "FL")
    tx_n = sum(1 for o in operators if o.get("socSite") == "TX")
    unk_n = sum(1 for o in operators if not o.get("socSite"))
    print(f"Imported {len(operators)} operators → {OUT}")
    print(f"  SOC sites: FL={fl_n}  TX={tx_n}  unknown={unk_n}")


if __name__ == "__main__":
    main()
